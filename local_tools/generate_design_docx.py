#!/usr/bin/env python3
"""Generate the competition design DOCX from the Markdown source."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


INLINE_CODE = re.compile(r"(`[^`]+`)")
HEADING = re.compile(r"^(#{1,6})\s+(.+)$")
ORDERED = re.compile(r"^(\d+)\.\s+(.+)$")
BULLET = re.compile(r"^-\s+(.+)$")
TABLE_SEPARATOR = re.compile(r"^\s*\|?(?:\s*:?-+:?\s*\|)+\s*:?-+:?\s*\|?\s*$")


def set_east_asia_font(run, name: str) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "打开 Microsoft Word 后更新目录域"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, placeholder, end])


def add_inline(paragraph, text: str, *, bold: bool = False) -> None:
    for part in INLINE_CODE.split(text):
        if not part:
            continue
        is_code = part.startswith("`") and part.endswith("`")
        content = part[1:-1] if is_code else part
        run = paragraph.add_run(content)
        run.bold = bold
        if is_code:
            set_east_asia_font(run, "Consolas")
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(150, 45, 45)
        else:
            set_east_asia_font(run, "宋体")


def split_table_row(line: str) -> list[str]:
    stripped = line.strip().strip("|")
    return [cell.strip() for cell in stripped.split("|")]


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Pt(21)

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(31, 55, 85)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    for style_name, size, color in (
        ("Heading 1", 16, RGBColor(31, 78, 121)),
        ("Heading 2", 14, RGBColor(47, 84, 150)),
        ("Heading 3", 12, RGBColor(54, 95, 145)),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    styles["Heading 1"].paragraph_format.page_break_before = True

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Consolas"
    code.font.size = Pt(8.5)
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    code.paragraph_format.first_line_indent = Pt(0)
    code.paragraph_format.left_indent = Cm(0.5)
    code.paragraph_format.right_indent = Cm(0.5)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(4)
    code.paragraph_format.line_spacing = 1.0


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.4)
    section.header_distance = Cm(1.3)
    section.footer_distance = Cm(1.3)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("SysY 编译器分析、设计与实现文档")
    set_east_asia_font(run, "宋体")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(117, 117, 117)
    add_page_number(section.footer.paragraphs[0])

    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    document.core_properties.title = "SysY 编译器分析、设计与实现文档"
    document.core_properties.subject = "曼巴奥特队 SysY 编译器设计文档"
    document.core_properties.author = "李东易"
    document.core_properties.keywords = "SysY, 编译器, RISC-V, SSA, MemorySSA"


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, row in enumerate(rows):
        for column_index in range(columns):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = row[column_index] if column_index < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, text, bold=row_index == 0)
            if row_index == 0:
                set_cell_shading(cell, "D9EAF7")
    set_repeat_table_header(table.rows[0])


def add_code_block(document: Document, code_lines: list[str]) -> None:
    paragraph = document.add_paragraph(style="Code Block")
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F3F5F7")
    p_pr.append(shading)
    run = paragraph.add_run()
    set_east_asia_font(run, "Consolas")
    for index, line in enumerate(code_lines):
        if index:
            run.add_break()
        run.add_text(line)


def render_markdown(document: Document, source: str) -> None:
    lines = source.splitlines()
    index = 0
    cover_table_done = False
    toc_inserted = False

    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue

        if line.startswith("```"):
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                code_lines.append(lines[index])
                index += 1
            index += 1
            add_code_block(document, code_lines)
            continue

        heading = HEADING.match(line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1:
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(paragraph, text, bold=True)
            else:
                if cover_table_done and not toc_inserted:
                    document.add_page_break()
                    toc_title = document.add_paragraph()
                    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    toc_run = toc_title.add_run("目录")
                    set_east_asia_font(toc_run, "黑体")
                    toc_run.bold = True
                    toc_run.font.size = Pt(18)
                    toc = document.add_paragraph()
                    toc.paragraph_format.first_line_indent = Pt(0)
                    add_toc(toc)
                    toc_inserted = True
                word_level = min(level - 1, 3)
                paragraph = document.add_paragraph(style=f"Heading {word_level}")
                add_inline(paragraph, text, bold=True)
            index += 1
            continue

        if line.lstrip().startswith("|"):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            rows = [split_table_row(item) for item in table_lines if not TABLE_SEPARATOR.match(item)]
            if rows:
                add_markdown_table(document, rows)
                if not cover_table_done:
                    cover_table_done = True
            continue

        ordered = ORDERED.match(line)
        if ordered:
            paragraph = document.add_paragraph(style="List Number")
            paragraph.paragraph_format.first_line_indent = Pt(0)
            add_inline(paragraph, ordered.group(2))
            index += 1
            continue

        bullet = BULLET.match(line)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.first_line_indent = Pt(0)
            add_inline(paragraph, bullet.group(1))
            index += 1
            continue

        paragraph_lines = [line.strip()]
        index += 1
        while index < len(lines):
            candidate = lines[index]
            if not candidate.strip():
                break
            if (candidate.startswith("```") or HEADING.match(candidate) or
                    candidate.lstrip().startswith("|") or ORDERED.match(candidate) or
                    BULLET.match(candidate)):
                break
            paragraph_lines.append(candidate.strip())
            index += 1
        paragraph = document.add_paragraph()
        add_inline(paragraph, " ".join(paragraph_lines))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    source = args.source.read_text(encoding="utf-8")
    document = Document()
    configure_styles(document)
    configure_document(document)
    render_markdown(document, source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)


if __name__ == "__main__":
    main()
